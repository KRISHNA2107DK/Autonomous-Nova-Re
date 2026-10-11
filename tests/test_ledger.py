"""Incident ledger: incidents, remediation chains, affected services and reasoning in one workbook + CSVs."""

from __future__ import annotations

import csv
import json

from novasurge import config, ledger
from novasurge.reasoning import Narrator


def _round() -> dict:
    return {
        "round": 1, "campaign_id": "c-test-ledger", "injector": "network_partition", "target": "order-service",
        "status": "RECOVERED", "injected_at": "2026-10-11T07:03:55+00:00", "detected_at": "2026-10-11T07:04:05+00:00",
        "mitigated_at": "2026-10-11T07:04:15+00:00", "recovered_at": "2026-10-11T07:04:55+00:00",
        "timeline": [{"at": "2026-10-11T07:03:36+00:00", "phase": "PREFLIGHT"}],
        "planner": {"decided_by": "llm", "model": "groq/test", "reason": "never tried on order-service"},
        "detection": {"affected_service": "order-service", "anomaly_type": "network_partition",
                      "evidence": ["pods unreachable"],
                      "anomalous_services": [{"service": "api-gateway", "severity": 0.9,
                                              "anomaly_type": "high_latency"}]},
        "rca": {"origin": "order-service", "anomaly_type": "network_partition", "confidence": 0.95,
                "call_path": ["api-gateway", "order-service"], "affected": ["api-gateway", "order-service"],
                "narrative": "Root cause localised to order-service.",
                "candidates": [{"service": "order-service", "score": 9.3, "reasons": ["earliest onset"]}]},
        "decision": {"action": "network_policy_rollback", "fallback": "rolling_restart",
                     "reasoning": ["partition -> roll back the policy"]},
        "remediations": [
            {"remediator": "network_policy_rollback", "service": "order-service", "success": True,
             "verified": False, "summary": "removed 0 policies", "duration_s": 2.0, "step": 1, "role": "primary",
             "reason": "partition -> roll back the policy", "started_at": "2026-10-11T07:04:10+00:00"},
            {"remediator": "rolling_restart", "service": "order-service", "success": True, "verified": True,
             "summary": "restarted 2 pods", "duration_s": 9.0, "step": 2, "role": "fallback",
             "reason": "rollback did not verify; escalating", "verified_at": "2026-10-11T07:04:15+00:00"},
        ],
        "healed_by": "fallback", "fault_cleared": True, "diagnosis_correct": True, "classification_correct": True,
        "sla": {"ttd_seconds": 10, "ttr_seconds": 50, "recovery_target_seconds": 90, "recovery_met": True,
                "detection_target_seconds": 30, "detection_met": True},
        "user_impact": {"total_requests": 1000, "failed_requests": 100, "failure_ratio": 0.1},
    }


def test_ledger_records_failure_chain_and_reasoning(sim_home):
    d = config.REPORTS_DIR / "c-test-ledger"
    d.mkdir(parents=True, exist_ok=True)
    (d / "rounds.json").write_text(json.dumps([_round()]))
    (d / "scorecard.json").write_text(json.dumps({"resilience_score": 99.0, "verdict": "EXCELLENT"}))
    n = Narrator("c-test-ledger", echo=False)
    n.round = 1
    n.say("REMEDIATE", "network_policy_rollback did not restore order-service; escalating")

    path = ledger.rebuild()
    assert path and path.exists()
    out = ledger.ledger_dir()

    with (out / "incidents.csv").open(encoding="utf-8-sig") as fh:
        inc = next(r for r in csv.DictReader(fh) if r["Incident ID"] == "c-test-ledger/R1")
    assert inc["Failure injected"] == "network_partition"
    assert inc["Injected at (UTC)"] == "2026-10-11 07:03:55"
    assert inc["Resolved by"] == "rolling_restart"
    assert "1. network_policy_rollback (FAILED)" in inc["Remediators applied (in order)"]
    assert "LLM agent" in inc["Chosen by"]

    with (out / "remediation_steps.csv").open(encoding="utf-8-sig") as fh:
        steps = [r for r in csv.DictReader(fh) if r["Incident ID"] == "c-test-ledger/R1"]
    assert [s["Remediator applied"] for s in steps] == ["network_policy_rollback", "rolling_restart"]
    assert steps[0]["Next remediator (if failed)"] == "rolling_restart"
    assert steps[0]["Outcome"].startswith("FAILED")
    assert steps[1]["Outcome"].startswith("RESOLVED")
    assert steps[1]["Role"] == "fallback"

    with (out / "affected_services.csv").open(encoding="utf-8-sig") as fh:
        roles = {r["Service"]: r["Role in incident"] for r in csv.DictReader(fh)
                 if r["Incident ID"] == "c-test-ledger/R1"}
    assert roles["order-service"] == "ROOT CAUSE"
    assert "victim" in roles["api-gateway"]

    with (out / "reasoning_timeline.csv").open(encoding="utf-8-sig") as fh:
        assert any("escalating" in r["Reasoning / message"] for r in csv.DictReader(fh))

    from openpyxl import load_workbook
    wb = load_workbook(path)
    assert wb.sheetnames == ["Incidents", "Remediation Steps", "Affected Services", "Reasoning Timeline",
                             "Campaigns", "About"]


def test_autopilot_incidents_are_recorded(sim_home):
    rec = _round() | {"incident_id": "AP-test-order-service", "started_at": "2026-10-11T09:00:00+00:00",
                      "finished_at": "2026-10-11T09:02:00+00:00", "planner": None}
    ledger.record_autopilot(rec)
    with (ledger.ledger_dir() / "incidents.csv").open(encoding="utf-8-sig") as fh:
        row = next(r for r in csv.DictReader(fh) if r["Incident ID"] == "AP-test-order-service")
    assert row["Source"] == "autopilot"
    assert row["Chosen by"] == "manual injection (CLI/API)"
