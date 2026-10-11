"""
Incident ledger: one durable, human-readable record of everything AutoHeal has done.

For every incident (a campaign round, or an anomaly healed by autopilot) it records
the failure injected, the services affected, the remedies suggested, every remediator
applied in order (and, when one failed, which one ran next and why), the final result,
the timestamps of each step and all the reasoning.

Written to <repo>/reports/ledger/:
  AUTOHEAL-INCIDENT-LEDGER.xlsx   one workbook, open it in Excel / LibreOffice / Google Sheets
      Incidents            one row per injected failure / incident, start to finish
      Remediation Steps    one row per remediator applied: why, result, verified, what ran next
      Affected Services    one row per service touched by an incident and its role
      Reasoning Timeline   every narration event, in order
      Campaigns            one row per campaign with its score
      About                what each sheet and column means
  incidents.csv, remediation_steps.csv, affected_services.csv, reasoning_timeline.csv, campaigns.csv
      the same tables as plain CSV (UTF-8, opens directly in Excel)

The ledger is regenerated from durable sources only — reports/<campaign>/rounds.json,
reports/<campaign>/events.jsonl and reports/autopilot/*.jsonl — so it is complete
and readable while AutoHeal is not running, and rebuilding it never loses data.
"""

from __future__ import annotations

import csv
import json
import logging
import os
import re
from datetime import datetime, timezone
from pathlib import Path

from novasurge import config, incident_log

log = logging.getLogger(__name__)

WORKBOOK = "AUTOHEAL-INCIDENT-LEDGER.xlsx"
AUTOPILOT_INCIDENTS = "incidents.jsonl"
_EVENT_FILE = re.compile(r"^\d{4}-\d{2}-\d{2}\.jsonl$")
_TEXT_LIMIT = 32000          # Excel cell limit is 32 767 characters


def ledger_dir() -> Path:
    return config.REPORTS_DIR / "ledger"


# --------------------------------------------------------------------------- helpers

def _ts(v) -> str:
    """ISO timestamp -> 'YYYY-MM-DD HH:MM:SS' (UTC), the form Excel reads as a date-time."""
    if not v:
        return ""
    try:
        dt = datetime.fromisoformat(str(v))
    except ValueError:
        return str(v)
    if dt.tzinfo:
        dt = dt.astimezone(timezone.utc)
    return dt.strftime("%Y-%m-%d %H:%M:%S")


def _yn(v) -> str:
    return "" if v is None else ("YES" if v else "NO")


def _num(v, nd: int = 2):
    return "" if v is None else (round(v, nd) if isinstance(v, (int, float)) else v)


def _join(items, sep: str = "\n") -> str:
    return sep.join(str(i) for i in (items or []) if i not in (None, ""))


def _status_summary(r: dict, steps: list[dict]) -> str:
    st = r.get("status")
    sla = r.get("sla") or {}
    ui = r.get("user_impact") or {}
    users = f"{ui['failure_ratio']:.1%} of user requests failed" if ui.get("failure_ratio") is not None else ""
    chain = " → ".join(f"{s['remediator']} ({'OK' if s['verified'] else 'failed'})" for s in steps)
    if st in ("RECOVERED", "MITIGATED", "MASKED"):
        last = next((s for s in reversed(steps) if s["verified"]), None)
        by = f"by {last['remediator']}" if last else f"via {r.get('healed_by') or 'self-healing'}"
        parts = [f"{st} {by}"]
        if sla.get("ttr_seconds") is not None:
            parts.append(f"recovered {sla['ttr_seconds']}s after detection "
                         f"(SLA {sla.get('recovery_target_seconds')}s {'met' if sla.get('recovery_met') else 'MISSED'})")
        elif st == "MITIGATED":
            parts.append("service verified working but the system did not return fully to normal in time")
        if st == "MASKED":
            parts.append("symptoms cleared but the injected fault was still present at teardown")
        if len(steps) > 1:
            parts.append(f"chain: {chain}")
        if users:
            parts.append(users)
        return "; ".join(parts)
    if st == "FAILED":
        return f"FAILED — every remediation option exhausted without verified recovery ({chain or 'none applied'})"
    if st == "UNDETECTED":
        return "UNDETECTED — no anomaly confirmed within the detection window; fault reverted at teardown"
    if st == "ABORTED":
        return "ABORTED by the user-impact watchdog; fault reverted immediately"
    if st == "SKIPPED_PREFLIGHT":
        return "Not injected — safety check said NO-GO: " + "; ".join((r.get("preflight") or {}).get("nogo_reasons") or [])
    if st == "SKIPPED":
        return "Not injected — no eligible (fault, service) pair"
    if st == "INJECTION_FAILED":
        return f"Injection failed: {(r.get('injection') or {}).get('error')}"
    return str(st or "")


# --------------------------------------------------------------------------- normalisation

def _steps(iid: str, r: dict) -> list[dict]:
    dec = r.get("decision") or {}
    rems = r.get("remediations") or []
    out = []
    for i, a in enumerate(rems):
        step = a.get("step") or i + 1
        role = a.get("role") or ("primary" if i == 0 else
                                 "fallback" if a.get("remediator") == dec.get("fallback") else "re-diagnosis")
        if a.get("reason"):
            why = a["reason"]
        elif i == 0:
            why = _join(dec.get("reasoning"), " ")
        else:
            prev = rems[i - 1].get("remediator")
            why = (f"{prev} did not verifiably restore the service; escalated to the playbook {role} "
                   f"{a.get('remediator')}.")
        diag = a.get("diagnosis") or {}
        diag_txt = (f"{diag.get('origin')} · {diag.get('anomaly_type')} · confidence {diag.get('confidence')}"
                    if diag else f"{(r.get('rca') or {}).get('origin')} · {(r.get('rca') or {}).get('anomaly_type')}")
        nxt = rems[i + 1].get("remediator") if i + 1 < len(rems) else ""
        verified = bool(a.get("verified"))
        if verified:
            outcome = "RESOLVED — service verified working (full Ready capacity + consecutive user-path probes OK)"
        elif not a.get("success"):
            outcome = "FAILED — the action itself did not complete"
        else:
            outcome = "FAILED — action completed but the service did not pass verification"
        if not verified:
            outcome += f"; next: {nxt}" if nxt else "; no further remediator tried"
        out.append({
            "incident_id": iid, "step": step,
            "started_at": _ts(a.get("started_at")), "completed_at": _ts(a.get("completed_at")),
            "service": a.get("service") or "", "diagnosis": diag_txt,
            "remediator": a.get("remediator") or "", "role": role, "why": why,
            "action_result": "completed" if a.get("success") else "failed",
            "verified": verified, "verified_at": _ts(a.get("verified_at")),
            "duration_s": _num(a.get("duration_s"), 1), "what_it_did": a.get("summary") or "",
            "details": json.dumps(a.get("details"), default=str) if a.get("details") else "",
            "outcome": outcome, "next_remediator": "" if verified else nxt,
        })
    return out


def _affected(iid: str, r: dict) -> list[dict]:
    rca = r.get("rca") or {}
    det = r.get("detection") or {}
    origin = rca.get("origin")
    sev = {s.get("service"): s for s in det.get("anomalous_services") or []}
    cands = {c.get("service"): c for c in rca.get("candidates") or [] if isinstance(c, dict)}
    services = list(dict.fromkeys([*([origin] if origin else []), *(rca.get("affected") or []), *sev, *cands]))
    call_path = rca.get("call_path") or []
    rows = []
    for s in services:
        if not s:
            continue
        if s == origin:
            role = "ROOT CAUSE"
        elif s in call_path:
            role = "caller on the failing path (victim)"
        else:
            role = "affected / deviating"
        sv = sev.get(s) or {}
        c = cands.get(s) or {}
        reasons = c.get("reasons")
        rows.append({
            "incident_id": iid, "service": s, "role": role,
            "detected_as": sv.get("anomaly_type") or "", "severity": _num(sv.get("severity"), 3),
            "first_flagged_at": _ts(sv.get("first_flagged_at")),
            "rca_score": _num(c.get("score"), 3),
            "rca_reasoning": _join(reasons, "; ") if isinstance(reasons, list) else str(reasons or ""),
        })
    return rows


def _incident(iid: str, r: dict, source: str, campaign: str, steps: list[dict]) -> dict:
    p = r.get("planner") or {}
    pf = r.get("preflight") or {}
    inj = r.get("injection") or {}
    det = r.get("detection") or {}
    rca = r.get("rca") or {}
    dec = r.get("decision") or {}
    sla = r.get("sla") or {}
    ui = r.get("user_impact") or {}
    if p:
        chosen_by = f"LLM agent ({p.get('model')})" if p.get("decided_by") == "llm" else "rule-based planner"
        why = p.get("reason") or ""
        if p.get("fallback_reason"):
            why += f" [rules used because: {p['fallback_reason']}]"
    elif source == "autopilot":
        chosen_by = ("manual injection (CLI/API)" if inj or r.get("injector")
                     else "not injected — organic anomaly")
        why = r.get("note") or ""
    else:
        chosen_by, why = "fixed campaign plan", ""
    affected = [a["service"] for a in _affected(iid, r)]
    verified = [s for s in steps if s["verified"]]
    return {
        "incident_id": iid, "source": source, "campaign": campaign, "round": r.get("round") or "",
        "final_status": r.get("status") or "", "result_summary": _status_summary(r, steps),
        "fault_injected": r.get("injector") or inj.get("injector") or "none",
        "target_service": r.get("target") or inj.get("service") or det.get("affected_service") or "",
        "injected_at": _ts(r.get("injected_at") or inj.get("injected_at")),
        "injection_details": json.dumps(inj.get("details"), default=str) if inj.get("details") else "",
        "chosen_by": chosen_by, "why_this_fault": why,
        "preflight": pf.get("go_nogo") or "", "injection_risk": pf.get("injection_risk") or "",
        "users_at_risk_pct": _num(pf.get("estimated_affected_users_pct"), 0),
        "detected_at": _ts(r.get("detected_at") or det.get("detected_at")),
        "ttd_s": _num(sla.get("ttd_seconds")), "detection_sla_s": _num(sla.get("detection_target_seconds"), 0),
        "detection_sla_met": _yn(sla.get("detection_met")) if r.get("injected_at") else "",
        "detected_as": det.get("anomaly_type") or "", "severity": _num(det.get("severity_score"), 3),
        "iforest_score": _num(det.get("iforest_score"), 4),
        "detection_evidence": _join(det.get("evidence")),
        "services_affected": ", ".join(affected),
        "root_cause": rca.get("origin") or "", "root_cause_type": rca.get("anomaly_type") or "",
        "rca_confidence": _num(rca.get("confidence"), 2), "rca_reasoning": rca.get("narrative") or "",
        "diagnosis_correct": _yn(r.get("diagnosis_correct")) if "diagnosis_correct" in r else "",
        "classification_correct": _yn(r.get("classification_correct")) if "classification_correct" in r else "",
        "remedy_suggested": dec.get("action") or "", "fallback_remedy": dec.get("fallback") or "",
        "decision_confidence": _num(dec.get("confidence"), 3), "decision_reasoning": _join(dec.get("reasoning")),
        "guardrails": _join(f"{g.get('id')} {g.get('outcome')}: {g.get('detail')}" if isinstance(g, dict) else g
                            for g in r.get("guardrails") or []),
        "remediators_applied": " → ".join(f"{s['step']}. {s['remediator']} ({'OK' if s['verified'] else 'FAILED'})"
                                          for s in steps),
        "attempts": len(steps), "resolved_by": verified[-1]["remediator"] if verified else "",
        "healed_via": r.get("healed_by") or "",
        "mitigated_at": _ts(r.get("mitigated_at")), "ttm_s": _num(sla.get("ttm_seconds")),
        "recovered_at": _ts(r.get("recovered_at")), "ttr_s": _num(sla.get("ttr_seconds")),
        "recovery_sla_s": _num(sla.get("recovery_target_seconds"), 0),
        "recovery_sla_met": _yn(sla.get("recovery_met")) if r.get("detected_at") else "",
        "fault_cleared": _yn(r.get("fault_cleared")) if "fault_cleared" in r else "",
        "user_requests": _num(ui.get("total_requests"), 0), "user_requests_failed": _num(ui.get("failed_requests"), 0),
        "user_failure_pct": _num(ui["failure_ratio"] * 100, 2) if ui.get("failure_ratio") is not None else "",
        "user_p99_ms": _num(ui.get("p99_ms"), 0),
    }


# --------------------------------------------------------------------------- sources

def _load_json(path: Path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return default


def _read_jsonl(path: Path) -> list[dict]:
    out = []
    try:
        for line in path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                try:
                    out.append(json.loads(line))
                except json.JSONDecodeError:
                    continue
    except OSError:
        pass
    return out


def collect() -> dict[str, list[dict]]:
    """Every table of the ledger, from the durable files under reports/."""
    tables: dict[str, list[dict]] = {k: [] for k in
                                     ("incidents", "steps", "affected", "timeline", "campaigns")}
    root = config.REPORTS_DIR
    if not root.exists():
        return tables
    for d in sorted(p for p in root.iterdir() if p.is_dir() and (p / "rounds.json").exists()):
        cid = d.name
        rounds = _load_json(d / "rounds.json", [])
        card = _load_json(d / "scorecard.json", {})
        for r in rounds:
            iid = f"{cid}/R{r.get('round')}"
            steps = _steps(iid, r)
            tables["steps"] += steps
            tables["affected"] += _affected(iid, r)
            tables["incidents"].append(_incident(iid, r, "campaign", cid, steps))
        for e in incident_log.load_events(cid):
            tables["timeline"].append(_event_row(e, f"{cid}/R{e['round']}" if e.get("round") else cid, cid))
        starts = [t["at"] for r in rounds for t in r.get("timeline") or []]
        comps = card.get("components") or {}
        tables["campaigns"].append({
            "campaign": cid, "started_at": _ts(min(starts)) if starts else "",
            "finished_at": _ts(card.get("generated_at")), "rounds": len(rounds),
            "resilience_score": card.get("resilience_score", ""), "verdict": card.get("verdict", ""),
            "recovery_sla_pts": comps.get("recovery_sla", ""), "detection_sla_pts": comps.get("detection_sla", ""),
            "effectiveness_pts": comps.get("remediation_effectiveness", ""),
            "diagnosis_pts": comps.get("diagnosis_accuracy", ""), "user_protection_pts": comps.get("user_protection", ""),
            "mean_ttd_s": card.get("mean_ttd_seconds", ""), "mean_ttr_s": card.get("mean_ttr_seconds", ""),
            "rca_accuracy": card.get("rca_origin_accuracy", ""),
            "classification_accuracy": card.get("classification_accuracy", ""),
            "mean_user_failure_pct": _num((card.get("mean_user_failure_ratio") or 0) * 100, 2)
            if card.get("mean_user_failure_ratio") is not None else "",
        })

    ap = root / "autopilot"
    if ap.exists():
        for r in _read_jsonl(ap / AUTOPILOT_INCIDENTS):
            iid = r["incident_id"]
            steps = _steps(iid, r)
            tables["steps"] += steps
            tables["affected"] += _affected(iid, r)
            tables["incidents"].append(_incident(iid, r, "autopilot", "", steps))
        windows = [(r["incident_id"], r.get("started_at") or "", r.get("finished_at") or "")
                   for r in _read_jsonl(ap / AUTOPILOT_INCIDENTS)]
        for f in sorted(p for p in ap.iterdir() if _EVENT_FILE.match(p.name)):
            for e in _read_jsonl(f):
                at = e.get("at") or ""
                iid = next((i for i, a, b in windows if a and b and a <= at <= b), "")
                tables["timeline"].append(_event_row(e, iid, ""))

    tables["incidents"].sort(key=lambda x: x["injected_at"] or x["detected_at"])
    tables["timeline"].sort(key=lambda x: x["timestamp"])
    return tables


def _event_row(e: dict, iid: str, cid: str) -> dict:
    return {"timestamp": _ts(e.get("at")), "incident_id": iid, "campaign": cid, "round": e.get("round") or "",
            "phase": e.get("phase") or "", "service": e.get("service") or "", "message": e.get("message") or ""}


# --------------------------------------------------------------------------- autopilot records

def record_autopilot(rec: dict) -> None:
    """Persist one autopilot incident (never raises) and refresh the ledger."""
    try:
        path = config.REPORTS_DIR / "autopilot" / AUTOPILOT_INCIDENTS
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("a", encoding="utf-8") as fh:
            fh.write(json.dumps(rec, default=str, ensure_ascii=False) + "\n")
    except OSError as exc:
        log.warning("could not record autopilot incident: %s", exc)
    rebuild()


# --------------------------------------------------------------------------- output

# (key, header, width, wrap)
COLUMNS: dict[str, list[tuple[str, str, int, bool]]] = {
    "incidents": [
        ("incident_id", "Incident ID", 30, False), ("source", "Source", 11, False),
        ("campaign", "Campaign", 24, False), ("round", "Round", 7, False),
        ("final_status", "Final status", 14, False), ("result_summary", "Final result", 60, True),
        ("fault_injected", "Failure injected", 19, False), ("target_service", "Target service", 18, False),
        ("injected_at", "Injected at (UTC)", 20, False), ("injection_details", "Injection details", 40, True),
        ("chosen_by", "Chosen by", 26, True), ("why_this_fault", "Why this failure / service", 60, True),
        ("preflight", "Safety check", 12, False), ("injection_risk", "Injection risk", 12, False),
        ("users_at_risk_pct", "Users at risk %", 10, False),
        ("detected_at", "Detected at (UTC)", 20, False), ("ttd_s", "Time to detect (s)", 11, False),
        ("detection_sla_s", "Detection SLA (s)", 11, False), ("detection_sla_met", "Detection SLA met", 11, False),
        ("detected_as", "Detected as", 18, False), ("severity", "Severity", 9, False),
        ("iforest_score", "Isolation Forest score", 11, False),
        ("detection_evidence", "Detection evidence", 60, True),
        ("services_affected", "Services affected", 34, True),
        ("root_cause", "Root cause (RCA)", 18, False), ("root_cause_type", "Root-cause type", 18, False),
        ("rca_confidence", "RCA confidence", 10, False), ("rca_reasoning", "RCA reasoning", 70, True),
        ("diagnosis_correct", "Root cause correct", 10, False),
        ("classification_correct", "Fault type correct", 10, False),
        ("remedy_suggested", "Remedy suggested (primary)", 22, False),
        ("fallback_remedy", "Fallback remedy", 18, False), ("decision_confidence", "Decision confidence", 10, False),
        ("decision_reasoning", "Decision reasoning", 70, True), ("guardrails", "Guardrails triggered", 40, True),
        ("remediators_applied", "Remediators applied (in order)", 44, True), ("attempts", "Attempts", 9, False),
        ("resolved_by", "Resolved by", 22, False), ("healed_via", "Healed via", 12, False),
        ("mitigated_at", "Mitigated at (UTC)", 20, False), ("ttm_s", "Time to mitigate (s)", 11, False),
        ("recovered_at", "Recovered at (UTC)", 20, False), ("ttr_s", "Time to recover (s)", 11, False),
        ("recovery_sla_s", "Recovery SLA (s)", 10, False), ("recovery_sla_met", "Recovery SLA met", 10, False),
        ("fault_cleared", "Fault gone at teardown", 11, False),
        ("user_requests", "User requests", 10, False), ("user_requests_failed", "User requests failed", 10, False),
        ("user_failure_pct", "User failure %", 10, False), ("user_p99_ms", "User p99 (ms)", 10, False),
    ],
    "steps": [
        ("incident_id", "Incident ID", 30, False), ("step", "Step", 6, False),
        ("started_at", "Started (UTC)", 20, False), ("completed_at", "Completed (UTC)", 20, False),
        ("service", "Service", 18, False), ("diagnosis", "Diagnosis at this step", 36, True),
        ("remediator", "Remediator applied", 22, False), ("role", "Role", 13, False),
        ("why", "Why this remediator", 70, True), ("action_result", "Action result", 11, False),
        ("verified", "Verified working", 10, False), ("verified_at", "Verified at (UTC)", 20, False),
        ("duration_s", "Duration (s)", 10, False), ("what_it_did", "What it did", 50, True),
        ("details", "Details", 40, True), ("outcome", "Outcome", 50, True),
        ("next_remediator", "Next remediator (if failed)", 22, False),
    ],
    "affected": [
        ("incident_id", "Incident ID", 30, False), ("service", "Service", 20, False),
        ("role", "Role in incident", 30, True), ("detected_as", "Detected as", 18, False),
        ("severity", "Severity", 9, False), ("first_flagged_at", "First flagged (UTC)", 20, False),
        ("rca_score", "RCA score", 9, False), ("rca_reasoning", "RCA reasoning", 90, True),
    ],
    "timeline": [
        ("timestamp", "Time (UTC)", 20, False), ("incident_id", "Incident ID", 30, False),
        ("campaign", "Campaign", 24, False), ("round", "Round", 7, False), ("phase", "Phase", 12, False),
        ("service", "Service", 18, False), ("message", "Reasoning / message", 120, True),
    ],
    "campaigns": [
        ("campaign", "Campaign", 26, False), ("started_at", "Started (UTC)", 20, False),
        ("finished_at", "Finished (UTC)", 20, False), ("rounds", "Rounds", 8, False),
        ("resilience_score", "Resilience score", 11, False), ("verdict", "Verdict", 50, True),
        ("recovery_sla_pts", "Recovery SLA /35", 10, False), ("detection_sla_pts", "Detection SLA /20", 10, False),
        ("effectiveness_pts", "Effectiveness /20", 10, False), ("diagnosis_pts", "Diagnosis /15", 10, False),
        ("user_protection_pts", "User protection /10", 10, False), ("mean_ttd_s", "Mean TTD (s)", 10, False),
        ("mean_ttr_s", "Mean TTR (s)", 10, False), ("rca_accuracy", "RCA accuracy", 10, False),
        ("classification_accuracy", "Classification accuracy", 12, False),
        ("mean_user_failure_pct", "Mean user failure %", 11, False),
    ],
}
SHEETS = {"incidents": ("Incidents", "incidents.csv"), "steps": ("Remediation Steps", "remediation_steps.csv"),
          "affected": ("Affected Services", "affected_services.csv"),
          "timeline": ("Reasoning Timeline", "reasoning_timeline.csv"), "campaigns": ("Campaigns", "campaigns.csv")}

ABOUT = [
    ("AutoHeal incident ledger", ""),
    ("Generated (UTC)", ""),
    ("", ""),
    ("Sheet", "What it holds"),
    ("Incidents", "One row per injected failure or autopilot incident: what was injected, why, when it was detected, "
                  "the services affected, the root cause and its reasoning, the remedy suggested, every remediator "
                  "applied in order, the final result, SLA and user impact."),
    ("Remediation Steps", "One row per remediator applied: the diagnosis at that moment, why it was chosen "
                          "(primary / fallback / re-diagnosis), whether it completed and was verified working, what "
                          "it did, and — if it failed — which remediator ran next."),
    ("Affected Services", "Every service touched by an incident: root cause, caller on the failing path, or other "
                          "deviating service, with detection severity and the RCA score and reasoning."),
    ("Reasoning Timeline", "Every plain-English narration event (plan, safety check, inject, detect, RCA, decide, "
                           "guardrail, remediate, verify, recover, teardown) in time order."),
    ("Campaigns", "One row per campaign with its resilience score and score breakdown."),
    ("", ""),
    ("Times", "All timestamps are UTC (India Standard Time = UTC + 5:30)."),
    ("Statuses", "RECOVERED = fixed and the whole system back to normal · MITIGATED = service verified working but "
                 "not fully back to normal in time · MASKED = symptoms gone but the fault was still present · "
                 "FAILED = every option exhausted · UNDETECTED = not detected in time · ABORTED = stopped by the "
                 "user-impact watchdog · SKIPPED_PREFLIGHT = not injected because the safety check said NO-GO."),
    ("Rebuild", "python -m novasurge ledger   (also refreshed automatically after every campaign and autopilot heal)"),
    ("Sources", "reports/<campaign>/rounds.json, reports/<campaign>/events.jsonl, reports/autopilot/*.jsonl"),
]


def _write_csvs(tables: dict[str, list[dict]], out: Path) -> None:
    for key, (_, fname) in SHEETS.items():
        cols = COLUMNS[key]
        with (out / fname).open("w", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh)
            w.writerow([h for _, h, _, _ in cols])
            for row in tables[key]:
                w.writerow([_cell_text(row.get(k, "")) for k, _, _, _ in cols])


def _cell_text(v):
    if isinstance(v, bool):
        return "YES" if v else "NO"
    if isinstance(v, str) and len(v) > _TEXT_LIMIT:
        return v[:_TEXT_LIMIT] + " …[truncated]"
    return v


def _write_workbook(tables: dict[str, list[dict]], path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="16213A")
    head_font = Font(bold=True, color="F6F7F9")
    thin = Side(style="thin", color="D3DAE3")
    border = Border(bottom=thin)
    fills = {"good": PatternFill("solid", fgColor="E3F1E7"), "warn": PatternFill("solid", fgColor="FBF0DC"),
             "bad": PatternFill("solid", fgColor="F8E1E1")}

    def tone(v: str) -> str | None:
        v = str(v)
        if v.startswith(("RECOVERED", "RESOLVED")) or v == "YES":
            return "good"
        if v.startswith(("MITIGATED", "MASKED", "SKIPPED")):
            return "warn"
        if v.startswith(("FAILED", "UNDETECTED", "ABORTED", "INJECTION_FAILED")) or v == "NO":
            return "bad"
        return None

    toned = {"final_status", "outcome", "verified", "detection_sla_met", "recovery_sla_met",
             "diagnosis_correct", "classification_correct", "fault_cleared"}
    times = {"injected_at", "detected_at", "mitigated_at", "recovered_at", "started_at", "completed_at",
             "verified_at", "first_flagged_at", "timestamp", "finished_at"}

    wb = Workbook()
    wb.remove(wb.active)
    for key, (title, _) in SHEETS.items():
        ws = wb.create_sheet(title)
        cols = COLUMNS[key]
        ws.append([h for _, h, _, _ in cols])
        for c, (_, _, width, _) in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=c)
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.row_dimensions[1].height = 34
        for row in tables[key]:
            values = []
            for k, _, _, _ in cols:
                v = _cell_text(row.get(k, ""))
                if k in times and v:
                    try:
                        v = datetime.strptime(v, "%Y-%m-%d %H:%M:%S")
                    except ValueError:
                        pass
                elif k == "verified" and isinstance(row.get(k), bool):
                    v = "YES" if row[k] else "NO"
                if isinstance(v, str):
                    v = ILLEGAL_CHARACTERS_RE.sub("", v)
                values.append(v)
            ws.append(values)
            r = ws.max_row
            for c, (k, _, _, wrap) in enumerate(cols, start=1):
                cell = ws.cell(row=r, column=c)
                cell.alignment = Alignment(wrap_text=wrap, vertical="top")
                cell.border = border
                if k in times and isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                if k in toned and (t := tone(cell.value)):
                    cell.fill = fills[t]
        ws.freeze_panes = "B2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("About")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 120
    for a, b in ABOUT:
        ws.append([a, b])
    ws["B2"] = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    ws["A1"].font = Font(bold=True, size=14)
    for row in ws.iter_rows(min_row=4):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    try:
        os.replace(tmp, path)
        return path
    except PermissionError:
        # The workbook is open in Excel (Windows locks it): keep the update under a new name.
        alt = path.with_name(f"{path.stem}-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx")
        os.replace(tmp, alt)
        log.warning("%s is open in another program; wrote %s instead", path.name, alt.name)
        return alt


def rebuild() -> Path | None:
    """Regenerate the whole ledger (workbook + CSVs). Never raises; returns the workbook path."""
    try:
        tables = collect()
        out = ledger_dir()
        out.mkdir(parents=True, exist_ok=True)
        try:
            _write_csvs(tables, out)
        except PermissionError as exc:
            log.warning("a ledger CSV is open in another program and was not updated: %s", exc)
        try:
            return _write_workbook(tables, out / WORKBOOK)
        except ImportError:
            log.warning("openpyxl is not installed; wrote the CSV ledger only")
            return out / "incidents.csv"
    except Exception as exc:  # noqa: BLE001 - the ledger must never break a campaign or a heal
        log.warning("could not rebuild the incident ledger: %s", exc)
        return None
